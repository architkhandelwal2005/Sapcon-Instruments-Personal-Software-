"""One-time seed import of the uncle's visit contact spreadsheet into
`entities` (+ an 'employer' relation linking each person to their company).

Clean fields only - company, person name, phone, email, region (from the zone
code), first_seen. NO title parsing: the sheet mashes the person's name and
designation together with a dash, so the name is taken heuristically (drop the
honorific, cut at the first dash/paren) and everything lands
`review_status='pending'` for the office boy to batch-confirm and fix in
/review. Near-duplicate companies in the sheet stay separate and are
reconciled later via the review queue.

Dry-run by default. --commit writes. --include-expo also pulls the expo sheet
(dates there are messy ranges, so first_seen is best-effort).
"""

import argparse
import re
import sys
import uuid
from datetime import date, datetime
from pathlib import Path
from typing import Optional

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from dotenv import load_dotenv

load_dotenv()

import openpyxl

from app.db import get_connection

XLSX = Path(__file__).resolve().parent.parent / "documents" / "Visit Scoop+Exhibition call List Updated (5).xlsx"

ZONE_MAP = {"MH": "MH", "ROI": "ROI", "NZ": "NZ", "SZ": "SZ", "GJ": "GJ", "EXPORT": "Export"}
FALLBACK_DATE = date(2025, 1, 1)

_HONORIFIC = re.compile(r"^(m/s|mr|mrs|ms|mst|dr|shri|smt|prof|capt|col)\b\.?\s*", re.I)
_NAME_CUT = re.compile(r"\s*[-–—(/]\s*")


def _clean(v) -> str:
    return "" if v is None else str(v).strip()


def _norm_key(s: str) -> str:
    return re.sub(r"\s+", " ", s.lower()).strip(" .,-")


def person_name(raw: str) -> Optional[str]:
    """"Mr. Anil Satwani-Managing Director" -> "Anil Satwani". The designation
    after the dash is dropped on purpose - it's the mashed field we don't
    trust."""
    s = _clean(raw)
    if not s or s in ("-", "--"):
        return None
    s = _HONORIFIC.sub("", s).strip()
    s = _NAME_CUT.split(s, maxsplit=1)[0].strip()
    s = re.sub(r"\s+", " ", s)
    return s or None


def company_name(raw: str) -> Optional[str]:
    s = re.sub(r"\s+", " ", _clean(raw)).strip(" -–—")
    return s or None


def first_token(raw: str) -> Optional[str]:
    s = _clean(raw)
    for sep in (",", "/", ";", "|"):
        s = s.split(sep)[0]
    return s.strip() or None


def parse_expo_date(raw: str) -> Optional[date]:
    s = _clean(raw)
    m = re.search(r"(\d{1,2})\s*[-–]\s*(\d{1,2})[/-](\d{1,2})[/-](\d{4})", s)
    if m:  # "03-05/01/2025" -> the 5th
        d, _, mth, yr = m.groups()
        try:
            return date(int(yr), int(mth), int(m.group(2)))
        except ValueError:
            return None
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def read_rows(include_expo: bool) -> list[dict]:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    out: list[dict] = []

    for r in wb["Visit Scoop Contacts"].iter_rows(min_row=3, values_only=True):
        company, name = company_name(r[3]), person_name(r[4])
        if not company and not name:
            continue
        out.append(
            {
                "company": company,
                "name": name,
                "phone": first_token(r[5]),
                "email": (first_token(r[6]) or "").lower() or None,
                "region": ZONE_MAP.get(_clean(r[2]).upper()),
                "first_seen": r[0].date() if hasattr(r[0], "date") else None,
                "source": "visit_list",
            }
        )

    if include_expo:
        for r in wb["Expo Contacts"].iter_rows(min_row=3, values_only=True):
            company, name = company_name(r[4]), person_name(r[5])
            if not company and not name:
                continue
            out.append(
                {
                    "company": company,
                    "name": name,
                    "phone": first_token(r[6]),
                    "email": (first_token(r[7]) or "").lower() or None,
                    "region": ZONE_MAP.get(_clean(r[3]).upper()),
                    "first_seen": parse_expo_date(r[1]),
                    "source": f"expo: {_clean(r[0])}" if _clean(r[0]) else "expo",
                }
            )
    return out


def collate(rows: list[dict]) -> tuple[dict, dict]:
    """Distinct company string -> one company; distinct (person, company) ->
    one person. earliest first_seen wins."""
    companies: dict[str, dict] = {}
    people: dict[tuple, dict] = {}
    for r in rows:
        if r["company"]:
            c = companies.setdefault(
                _norm_key(r["company"]),
                {"canonical_name": r["company"], "region": r["region"], "source": r["source"], "first_seen": r["first_seen"]},
            )
            if r["first_seen"] and (not c["first_seen"] or r["first_seen"] < c["first_seen"]):
                c["first_seen"] = r["first_seen"]
        if r["name"]:
            key = (_norm_key(r["name"]), _norm_key(r["company"] or ""))
            pe = people.setdefault(
                key,
                {"canonical_name": r["name"], "company": r["company"],
                 **{k: r[k] for k in ("phone", "email", "region", "source", "first_seen")}},
            )
            if r["first_seen"] and (not pe["first_seen"] or r["first_seen"] < pe["first_seen"]):
                pe["first_seen"] = r["first_seen"]
    return companies, people


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--include-expo", action="store_true", help="also import the expo sheet")
    ap.add_argument("--commit", action="store_true", help="write to the DB (default: dry-run)")
    args = ap.parse_args()

    rows = read_rows(args.include_expo)
    companies, people = collate(rows)
    company_only = sum(1 for r in rows if r["company"] and not r["name"])
    print(f"Read {len(rows)} rows -> {len(companies)} distinct companies, {len(people)} distinct people "
          f"({company_only} rows were company-only).")
    print("\nSample people:")
    for pe in list(people.values())[:12]:
        print(f"  {pe['canonical_name']!r:26} @ {(pe['company'] or '-')[:34]:34} | {pe['region'] or '-':6} | {pe['first_seen']}")

    if not args.commit:
        print("\nDRY RUN - nothing written. Re-run with --commit.")
        return

    # Pre-generate ids client-side and bulk-insert with three executemany calls
    # rather than ~1200 RETURNING round-trips - the pooled connection drops
    # under that many small statements.
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("select lower(canonical_name), entity_type, id from entities")
            existing = {(n, t): str(i) for n, t, i in cur.fetchall()}

        company_ids: dict[str, str] = {}
        company_rows = []
        for key, c in companies.items():
            hit = existing.get((c["canonical_name"].lower(), "company"))
            if hit:
                company_ids[key] = hit
                continue
            cid = str(uuid.uuid4())
            company_ids[key] = cid
            company_rows.append((cid, c["canonical_name"], c["region"], c["source"], c["first_seen"]))

        people_rows, relation_rows = [], []
        for (nk, ck), pe in people.items():
            if (pe["canonical_name"].lower(), "person") in existing:
                continue
            pid = str(uuid.uuid4())
            people_rows.append((pid, pe["canonical_name"], pe["phone"], pe["email"], pe["region"], pe["source"], pe["first_seen"]))
            cid = company_ids.get(ck)
            if cid:
                relation_rows.append((str(uuid.uuid4()), pid, cid, pe["first_seen"] or FALLBACK_DATE))

        with conn.cursor() as cur:
            cur.executemany(
                "insert into entities (id, canonical_name, entity_type, region, source, first_seen, review_status) "
                "values (%s,%s,'company',%s,%s,%s,'pending')",
                company_rows,
            )
            cur.executemany(
                "insert into entities (id, canonical_name, entity_type, phone, email, region, source, first_seen, review_status) "
                "values (%s,%s,'person',%s,%s,%s,%s,%s,'pending')",
                people_rows,
            )
            cur.executemany(
                "insert into relations (id, source_id, target_id, role_tag, description, provenance, recorded_at, review_status) "
                "values (%s,%s,%s,'employer','works at','direct',%s,'pending')",
                relation_rows,
            )
        conn.commit()
        print(f"\nCommitted: {len(company_rows)} companies, {len(people_rows)} people, "
              f"{len(relation_rows)} employer relations - all pending review.")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
